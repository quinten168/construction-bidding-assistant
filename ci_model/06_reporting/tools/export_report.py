"""
Tool 5 — export_report

Writes the assembled report text to disk.
  - format="md"   → markdown file (always available)
  - format="docx" → Word document (requires python-docx)
"""
from __future__ import annotations

from pathlib import Path


def export_report(
    report_text: str,
    output_path: str | Path,
    fmt: str = "md",
) -> Path:
    """
    Write report_text to output_path.

    Args:
        report_text:  full report as a string
        output_path:  destination file path (extension added if missing)
        fmt:          "md" or "docx"

    Returns:
        Path to the written file.
    """
    output_path = Path(output_path)

    if fmt == "docx":
        try:
            from docx import Document
            from docx.shared import Pt
        except ImportError:
            print("  [export] python-docx not installed — falling back to markdown.")
            fmt = "md"

    if fmt == "docx":
        if output_path.suffix.lower() != ".docx":
            output_path = output_path.with_suffix(".docx")
        doc = Document()
        for line in report_text.splitlines():
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("━" * 5):
                doc.add_paragraph("─" * 60)
            else:
                doc.add_paragraph(line)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
    else:
        if output_path.suffix.lower() not in (".md", ".txt"):
            output_path = output_path.with_suffix(".md")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(report_text, encoding="utf-8")

    return output_path
