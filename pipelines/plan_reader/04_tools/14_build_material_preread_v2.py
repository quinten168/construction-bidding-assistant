#!/usr/bin/env python3
"""
14_build_material_preread_v2.py — the SHORT version of the material pre-read (md + docx).

WHY v2: v1 (12_ + 13_) lists every spec'd variant as its own row — 149 rows across the five columns.
That is the right level of detail for the record, and too much to review in a meeting. T asked for a
version with fewer types that is easier to sign off.

WHAT CHANGES from v1 — one idea: the ROW IS THE MATERIAL, not the variant.
  • one row per common material (Aggregate Surfacing), not per class (Class 1 / 5 / 2 / Special)
  • the classes are NOT deleted — they move into a "Classes seen" column on that row, so CSI still
    sees FAA 42/43/45 and CL 5/13 (T, 2026-07-30: "these are all really important"). The difference
    is that it is now ONE decision instead of nine.
  • materials seen in only ONE plan are dropped from the document entirely. They were a footnote
    through v3; T, 2026-08-03 asked for the footnotes gone. v1 remains the complete record.

If CSI rules that a class inside a material behaves differently (the open question "is Aggregate
Surfacing Class 5 base rather than shoulder?"), that material gets split out — the data already
supports it, this version just does not lead with it.

v3, 2026-08-03 (T's review):
  • the checkboxes are CLICKABLE in the .docx, not glyphs to print and pen in — see
    13_render_preread_docx.add_checkbox for why it is a w14:checkbox content control and not a
    legacy form field. Called out in the opening callout so CSI knows to click rather than print.
  • no em dashes anywhere in the client-facing text, in prose or in generated labels
  • the "Why this changed" section (Stearns, 1.6M SY uncounted) is INTERNAL and is not rendered.
    The Change column stays — its cell values read on their own.
  • the seven judgment-call questions are not rendered here either. They are kept below in
    QUESTIONS, deliberately, so they are not lost from the project — see the note on that constant.
  • riprap is flagged inside Mixed_tons rather than as a note under `other`, because "is riprap
    asphalt?" is a Mixed_tons question and nobody reads the footnotes of a bucket they skipped.

Run:
    python3 pipelines/plan_reader/04_tools/14_build_material_preread_v2.py
"""
import collections
import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

TOOLS_DIR = Path(__file__).resolve().parent


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, TOOLS_DIR / path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


pre = _load("preread", "12_build_material_preread.py")     # collect(), labels, bucket names
dx  = _load("prdocx", "13_render_preread_docx.py")         # the shared docx styling helpers

OUT_MD   = pre.OUT_DIR / "2026-08-03_material_mapping_for_CSI_v3.md"
OUT_DOCX = pre.OUT_DIR / "2026-08-03_material_mapping_for_CSI_v3.docx"

MIN_PLANS = 2          # a material must appear in at least this many plans to be shown at all
MAX_UNPLACED = 15      # this is a sample; v1 carries all 45

# The three rulings CSI makes on every material row. One list, used for both the markdown glyphs and
# the clickable docx controls, so the two outputs cannot say different things.
ROW_CHOICES     = ["Agree", "Move to ______", "Do not count"]
OTHER_CHOICES   = ["Agree, flag only", "Count as ______"]
COLUMN_CHOICES  = ["Mixed_tons", "base_tons", "shldr_tons", "grading_yards", "mill_sy",
                   "Do not count"]
RIPRAP_CHOICES  = ["Agree, riprap is not asphalt", "No, count riprap in Mixed_tons"]


def md_boxes(choices):
    return " ".join("☐ " + c for c in choices)


# T, 2026-08-03: "remove all em dashes". Prose in this file is written without them; this scrubs the
# ones that arrive from the DATA (pay-item names) or from shared helpers in 12_, which still uses an
# em dash as its empty-quantity placeholder. Applied at every point a derived string reaches the page,
# and re-checked at the end of main() so a new code path cannot leak one in silently.
DASHES = "—–−"


def nd(s):
    """No dashes: em dash, en dash and minus sign out of any string headed for the document."""
    s = str(s)
    for d in DASHES:
        s = s.replace(d, "-")
    return re.sub(r"\s*-\s*-\s*", " - ", s).strip()


def blank(s, placeholder=""):
    """12_ renders an absent quantity as an em dash. Here that is just empty."""
    s = nd(s)
    return placeholder if s in ("-", "") else s


def classes_seen(members):
    """Pull the class / grade designators off the printed names for the 'Classes seen' column.

    This is the inverse of base_name(): base_name strips the designator to GROUP, this keeps only the
    designator to DISPLAY. Nothing is invented — every token here was printed on a bid schedule.

    RAP is picked up alongside them (T, 2026-08-03). SUPERPAVE FAA 43 and RAP - SUPERPAVE FAA 43 are
    one material with a grade and a recycled-content flag, and base_name already groups them; without
    this the grouped row silently swallowed the fact that some of it is recycled. RAP is a flag, not
    a grade, so it is listed last rather than sorted in among the FAA numbers.
    """
    found, rap = [], False
    for name, _e in members:
        if re.search(r"\bRAP\b", name):
            rap = True
        for m in re.finditer(r"\b(CLASS|CL|FAA|TYPE SP|SP|GRADE|PG)\s*([A-Z0-9][A-Z0-9.\-]*)", name):
            kind, val = m.group(1), m.group(2)
            kind = {"CLASS": "Cl", "CL": "Cl", "TYPE SP": "Type SP", "SP": "SP"}.get(kind, kind)
            tok = f"{kind} {val.title() if val.isalpha() else val}"
            if tok not in found:
                found.append(tok)
    if rap:
        found.append("RAP")
    return nd(", ".join(found))


def n_plans(agg):
    """How many plans this document is actually based on. NOT the number of plans in the database.
    The DB holds 117, but only the plans with a priced bid schedule contributed a material line, and
    quoting 117 to CSI would overstate the evidence behind every "seen in N plans" figure."""
    return len({p for e in agg.values() for p in e["plans"]})


def change_label(today, proposed):
    """What this mapping CHANGES for the material, which is the only thing CSI needs from the old
    behaviour. v2 showed a raw "Today" value inside a section headed "What we ARE tracking", so a row
    could read "tracking" and "not counted" at once — contradictory (T, 2026-08-03). Framed as a
    change, the same fact becomes the argument for the mapping rather than a confusion:
        NEW        the quantity is invisible today and this mapping recovers it
        MOVED      it counts today, but in a different column
        no change  already correct
    The wording has to stand alone: the section that used to explain this column was the internal
    Stearns write-up, and that is no longer in the client document.
    """
    real = {t for t in today if t != "not counted"}
    missing = "not counted" in today
    if not real:
        return ("NEW, not counted today", True)
    if real == {proposed}:
        return ("no change", False) if not missing else ("partly NEW", True)
    other = ", ".join(sorted(real - {proposed}))
    return (f"MOVED from {other}", True)


def build_rows(bucket, agg):
    """(rows shown, one-off materials dropped) for one CSI column, most-frequent first."""
    items = [(n, e) for (b, n), e in agg.items() if b == bucket and e["qty"]]
    groups = collections.defaultdict(list)
    for name, e in items:
        groups[e["base"]].append((name, e))

    shown, tail = [], []
    for _key, members in groups.items():
        plans = {p for _n, e in members for p in e["plans"]}
        qty = collections.defaultdict(float)
        states, units, today = set(), set(), set()
        for _n, e in members:
            for u, v in e["qty"].items():
                qty[u] += v
            states |= e["states"]
            units |= e["units"]
            today |= e["today"]
        label = pre.base_label([n for n, _e in members])
        entry = {
            "label": nd(label), "classes": classes_seen(members), "plans": len(plans),
            "states": ",".join(sorted(s for s in states if s and s != "?")),
            "units": ",".join(sorted(units)),
            "change": change_label(today, pre.CSI_BUCKET[bucket]),
            "qty": "; ".join(f"{v:,.0f} {u}" for u, v in sorted(qty.items())),
            "sort": (-len(plans), -sum(qty.values())),
        }
        (shown if len(plans) >= MIN_PLANS else tail).append(entry)
    shown.sort(key=lambda r: r["sort"])
    tail.sort(key=lambda r: r["sort"])
    return shown, tail


# The riprap flag. It lives in Mixed_tons, not in the notes under `other`: the question is whether
# riprap is asphalt, and an estimator who agrees with the Mixed_tons table will never scroll to a
# footnote three sections later. T, 2026-08-03 flagged it as needing an explicit confirmation.
RIPRAP_NOTE = (
    "One to confirm: riprap. Riprap is quarried rock slope protection, placed at culvert ends and "
    "on slopes and measured in CY or TON. We do NOT count it as asphalt and it does not reach "
    "Mixed_tons; we carry it under structures as a scope flag. An early draft of our asphalt list "
    "had it in the wrong place, so we want this one confirmed rather than assumed."
)


def render_markdown(agg):
    out = []
    w = out.append
    w("# Material Mapping for CSI review")
    w("Proposed roll-up of plan materials into your quantity columns  ·  v3  ·  August 3, 2026\n")
    w("## What this is\n")
    w("This is **the mapping we intend to use**. We read every priced line item off the bid "
      "schedules in the plans we have processed and assigned each material to one of your six "
      "buckets. **If there are no objections, we will use this.**\n")
    w("Mark any row you disagree with: **Move** it to a different bucket, or **Do not count** it if "
      "it should not be counted at all. Most-frequent materials first.\n")
    w("**In the Word version every checkbox is clickable.** Open it in Word and click the boxes; you "
      "do not need to print it. The blanks and the empty tables are typed into the same way.\n")
    w(f"Based on **{n_plans(agg)} plans** with a priced bid schedule (MN / ND / SD).\n")
    w("## The six buckets\n")
    w("| Bucket | What goes in it |")
    w("|---|---|")
    w("| `Mixed_tons` | Hot mix asphalt placed, TONS |")
    w("| `base_tons` | Aggregate base placed under the pavement, TONS |")
    w("| `shldr_tons` | Shoulder gravel and aggregate surfacing, TONS |")
    w("| `grading_yards` | Earthwork moved, CUBIC YARDS |")
    w("| `mill_sy` | Existing surface ground off, SQUARE YARDS |")
    w("| `other` | Real work, but not one of your five quantities. Carried as a scope flag |")
    w("")
    w("---\n")
    w("# The mapping\n")
    for bucket in pre.TRACKED_ORDER:
        col = pre.CSI_BUCKET[bucket]
        shown, _tail = build_rows(bucket, agg)
        w(f"## `{col}`: {nd(pre.BUCKET_BLURB[bucket])}\n")
        w("| Ref | Material | Classes seen | Plans | States | Unit | Change | Object? |")
        w("|---|---|---|---|---|---|---|---|")
        for i, r in enumerate(shown, 1):
            txt, loud = r["change"]
            w(f"| {bucket[:2].upper()}{i} | **{r['label']}** | {r['classes']} | {r['plans']} | "
              f"{r['states']} | {r['units']} | {'**' + txt + '**' if loud else txt} | "
              f"{md_boxes(ROW_CHOICES)} |")
        w("")
        if bucket == "ASPHALT":
            w(f"> **{RIPRAP_NOTE}**  {md_boxes(RIPRAP_CHOICES)}\n")
        w(f"**Anything missing from `{col}`?**\n")
        w("| Material we did not list | Why it belongs here |")
        w("|---|---|")
        w("|  |  |\n")

    w("---\n")
    w("## `other`: real work, not one of your five quantities\n")
    w("These are real pay items. They do not roll into a quantity column. We carry them as a yes/no "
      "scope flag on the report instead. **Object if any of these should be counted.**\n")
    w("| Group | Plans | Example materials | Object? |")
    w("|---|---|---|---|")
    for label, g in not_tracked(agg):
        ex = "; ".join(nd(n.title()) for n, _c in sorted(g["names"], key=lambda x: -x[1])[:3])
        w(f"| **{nd(label)}** | {len(g['plans'])} | {ex} | {md_boxes(OTHER_CHOICES)} |")
    w("")
    w("- **Crushing and stockpile** are not extracted: filled on under 5% of your historical rows, "
      "and of the 117 plan sets we have processed, 1 had stockpile and none had crushing. Correct "
      "to treat these as internal sourcing decisions rather than plan quantities?\n")

    w("---\n")
    w("## Still unassigned: we need your call\n")
    unplaced = sorted([(k[1], v) for k, v in agg.items() if k[0] == "(UNPLACED)" and v["qty"]],
                      key=lambda kv: (-len(kv[1]["plans"]), -sum(kv[1]["qty"].values())))
    w(f"The {MAX_UNPLACED} most common of {len(unplaced)} priced materials we could not confidently "
      "assign. Most likely to be wrong, so worth your eyes.\n")
    w("| Material | Plans | States | Unit | Quantity | Which column? |")
    w("|---|---|---|---|---|---|")
    for name, e in unplaced[:MAX_UNPLACED]:
        w(f"| {nd(name.title())} | {len(e['plans'])} | "
          f"{','.join(sorted(s for s in e['states'] if s and s != '?'))} | "
          f"{','.join(sorted(e['units']))} | {blank(pre.fmt_qty(e))} | "
          f"{md_boxes(COLUMN_CHOICES)} |")
    w("")
    OUT_MD.write_text("\n".join(out))


# NOT RENDERED, AND KEPT ON PURPOSE.
#
# These are the seven calls we could not make for CSI. T, 2026-08-03 cut the section from the client
# document — the document is a proposal to be agreed to, and a page of open questions at the back
# invites the meeting to end unresolved instead. The questions themselves are still live and still
# unanswered, and each one is worth real quantity (FDR alone is 1.17M SY), so they stay here rather
# than being deleted: they are the agenda for the conversation, not a section of the handout.
# Do NOT let a "remove unused code" pass delete this.
QUESTIONS = [
    ("Full Depth Reclamation, 1,166,000 SY, 14 plans, currently counted as nothing.",
     "FDR grinds the asphalt and the base below it into a new base layer. Is that `mill_sy`, is it "
     "`grading_yards` because you are rebuilding the roadbed, or does it need its own column?"),
    ("Milling is measured in TONS in North Dakota, 230,625 T.",
     "MN and SD measure it in SY. `mill_sy` cannot hold both. Separate tons line, or a conversion?"),
    ("Topsoil, 539,000 CY.",
     "Currently filed as erosion control, so it never reaches `grading_yards`. Is topsoil part of "
     "your grading number?"),
    ("Pavement removal.",
     "Mark said it \"is considered grading work\" and converts to CY. (a) Only pavement and base "
     "removal, not remove-silt-fence / sign / pipe? (b) Is 1.75 Ton/CY your standard factor?"),
    ("Aggregate Surfacing Class 1 vs Class 5 vs Special, 67,000 T.",
     "We put all of it in `shldr_tons`. Is Class 5 really base material to you?"),
    ("Salvaged / remove-and-relay base.",
     "Re-used, not bought. Does it count toward `base_tons` because you still move and place it?"),
    ("PG binder, tack and prime, 56,029 T.",
     "We propose excluding these from `Mixed_tons` so mix tonnage is not double-counted. Confirm?"),
]

# Also internal, also kept: the Stearns evidence that motivated the whole remap. It is the argument
# for doing this, not something CSI needs to read, and T ruled it internal on 2026-08-03.
WHY_INTERNAL = (
    "On Stearns County CP 073-123-025 we read all 26 line items correctly but put them in the wrong "
    "buckets: 70,263 SY of milling counted as asphalt and then dropped, and 70,167 SY of "
    "reclamation not counted at all. The job reported zero milling on a job that is almost entirely "
    "milling and reclamation. Corpus-wide that was 1.6M SY of milling and 1.2M SY of reclamation "
    "going uncounted."
)


def not_tracked(agg):
    groups = collections.defaultdict(lambda: {"plans": set(), "names": [],
                                              "qty": collections.defaultdict(float)})
    for (bucket, name), e in agg.items():
        if bucket != "OTHER":
            continue
        label = pre.NOT_TRACKED_GROUPS.get(next(iter(e["family"]), None), "Other")
        g = groups[label]
        g["plans"] |= e["plans"]
        g["names"].append((name, len(e["plans"])))
        for u, v in e["qty"].items():
            g["qty"][u] += v
    return sorted(groups.items(), key=lambda kv: -len(kv[1]["plans"]))


def render_docx(agg):
    doc = Document()
    doc.styles["Normal"].font.name = dx.FONT
    doc.styles["Normal"].font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.6)
        s.top_margin = s.bottom_margin = Inches(0.6)

    dx.para(doc, "Material Mapping for CSI review", size=22, bold=True, color=dx.NAVY, space_after=2)
    dx.para(doc, "Proposed roll-up of plan materials into your quantity columns  ·  v3  ·  "
                 "August 3, 2026", size=10, color=dx.MUTED, space_after=12)
    dx.callout(doc, [
        ("This is the mapping we intend to use. If there are no objections, we will use it.", True),
        (f"We read every priced line item off the bid schedules in {n_plans(agg)} plans with a "
         "priced bid schedule (MN / ND / SD) and assigned each material to one of your six buckets.",
         False),
        ("Mark any row you disagree with: Move it to a different bucket, or Do not count it. "
         "Most-frequent materials first.", False),
        ("The checkboxes in this document are clickable. Open it in Word and click a box to tick "
         "it. No need to print it. Type into the blanks and the empty tables the same way.", True),
    ])
    dx.heading(doc, "The six buckets")
    dx.build_table(doc, ["Bucket", "What goes in it"],
                   [[("Mixed_tons", None, True), "Hot mix asphalt placed, TONS"],
                    [("base_tons", None, True), "Aggregate base placed under the pavement, TONS"],
                    [("shldr_tons", None, True), "Shoulder gravel and aggregate surfacing, TONS"],
                    [("grading_yards", None, True), "Earthwork moved, CUBIC YARDS"],
                    [("mill_sy", None, True), "Existing surface ground off, SQUARE YARDS"],
                    [("other", None, True),
                     "Real work, but not one of your five quantities. Carried as a scope flag"]],
                   [1.5, 7.0])

    dx.heading(doc, "The mapping")
    for bucket in pre.TRACKED_ORDER:
        col = pre.CSI_BUCKET[bucket]
        shown, _tail = build_rows(bucket, agg)
        dx.heading(doc, f"{col}:  {nd(pre.BUCKET_BLURB[bucket])}", level=2)
        dx.build_table(
            doc,
            ["Ref", "Material", "Classes seen", "Plans", "States", "Unit", "Change", "Object?"],
            [[f"{bucket[:2].upper()}{i}", (r["label"], None, True), r["classes"], str(r["plans"]),
              r["states"], r["units"],
              (r["change"][0], dx.RED, True) if r["change"][1] else r["change"][0],
              list(ROW_CHOICES)]
             for i, r in enumerate(shown, 1)],
            [0.4, 1.9, 1.5, 0.5, 0.6, 0.6, 1.5, 2.1])
        if bucket == "ASPHALT":
            dx.callout(doc, [(RIPRAP_NOTE, False), (list(RIPRAP_CHOICES), False)])
        dx.para(doc, f"Anything missing from {col}?", size=9.5, bold=True, space_after=3)
        dx.build_table(doc, ["Material we did not list", "Why it belongs here"],
                       [["", ""]], [4.3, 4.3])

    doc.add_page_break()
    dx.heading(doc, "other:  real work, not one of your five quantities", level=2)
    dx.para(doc, "These are real pay items. They do not roll into a quantity column. We carry them "
                 "as a yes/no scope flag on the report instead.", space_after=4)
    dx.para(doc, "Object if any of these should be counted.", bold=True, space_after=8)
    dx.build_table(doc, ["Group", "Plans", "Example materials", "Object?"],
                   [[(nd(label), None, True), str(len(g["plans"])),
                     "; ".join(nd(n.title())
                               for n, _c in sorted(g["names"], key=lambda x: -x[1])[:3]),
                     list(OTHER_CHOICES)]
                    for label, g in not_tracked(agg)],
                   [2.4, 0.5, 3.4, 2.3])
    dx.callout(doc, [
        ("One to call out", True),
        ("Crushing and stockpile are not extracted. They are filled on under 5% of your historical "
         "rows, and of the 117 plan sets we have processed, 1 had stockpile and none had crushing. "
         "Correct to treat these as internal sourcing decisions rather than plan quantities?",
         False),
    ])

    dx.heading(doc, "Still unassigned: we need your call", level=2)
    unplaced = sorted([(k[1], v) for k, v in agg.items() if k[0] == "(UNPLACED)" and v["qty"]],
                      key=lambda kv: (-len(kv[1]["plans"]), -sum(kv[1]["qty"].values())))
    dx.para(doc, f"The {MAX_UNPLACED} most common of {len(unplaced)} priced materials we could not "
                 "confidently assign. Most likely to be wrong, so worth your eyes.", space_after=8)
    dx.build_table(doc, ["Material", "Plans", "States", "Unit", "Quantity", "Which column?"],
                   [[nd(n.title()), str(len(e["plans"])),
                     ",".join(sorted(s for s in e["states"] if s and s != "?")),
                     ",".join(sorted(e["units"])), blank(pre.fmt_qty(e)),
                     list(COLUMN_CHOICES)]
                    for n, e in unplaced[:MAX_UNPLACED]],
                   [2.3, 0.5, 0.6, 0.6, 1.2, 3.4])
    doc.save(OUT_DOCX)


def audit_dashes():
    """Fail loudly rather than shipping an em dash. T asked for none; a generated document is the
    kind of place one creeps back in through a pay-item name nobody looked at."""
    from docx import Document as _D
    bad = [d for d in DASHES if d in OUT_MD.read_text()]
    doc = _D(str(OUT_DOCX))
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        texts += [c.text for row in t.rows for c in row.cells]
    bad += [d for d in DASHES if any(d in t for t in texts)]
    if bad:
        raise AssertionError(f"em/en dash still present in output: {sorted(set(bad))}")


def main():
    agg = pre.collect()
    render_markdown(agg)
    render_docx(agg)
    audit_dashes()
    total = sum(len(build_rows(b, agg)[0]) for b in pre.TRACKED_ORDER)
    dropped = sum(len(build_rows(b, agg)[1]) for b in pre.TRACKED_ORDER)
    print(f"v3 rows shown: {total}   one-off materials dropped: {dropped}")
    for b in pre.TRACKED_ORDER:
        shown, t = build_rows(b, agg)
        print(f"  {pre.CSI_BUCKET[b]:16s} {len(shown):>3} rows  (-{len(t)} one-offs)")
    print("no em/en dashes in either output")
    print(f"wrote {OUT_MD.name} + {OUT_DOCX.name}")


if __name__ == "__main__":
    main()
